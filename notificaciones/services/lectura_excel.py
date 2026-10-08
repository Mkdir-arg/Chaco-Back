"""Lectura del Excel de destinatarios de una campaña (RN-007-02 a RN-007-06, RN-007-15).

Sigue el patrón de ``programas.services.padron.parsear_padron``: ``.xlsx`` de hasta 2 MB,
techo de tamaño descomprimido leído del índice del zip (SEC-31), ``openpyxl`` en
``read_only`` y ``data_only`` (una fórmula se lee por su valor calculado) y
``ValidationError`` con un mensaje que el operador entiende.

El Excel trae **solo correos**: se toma la primera hoja y, de ella, la columna cuyo
encabezado es ``email``, ``correo`` o ``mail`` (sin distinguir mayúsculas ni acentos);
si la primera fila no tiene un encabezado reconocible, se lee la columna A desde la
fila 1. Cualquier otra columna se ignora.
"""

from __future__ import annotations

import unicodedata
import zipfile
from dataclasses import dataclass, field

from django.core.exceptions import ValidationError
from django.core.validators import EmailValidator

from notificaciones.models import TOPE_DESTINATARIOS, Descartado

EXCEL_MAX_BYTES = 2 * 1024 * 1024
#: Techo del contenido **descomprimido**: un .xlsx de 2 MB puede inflarse a cientos.
EXCEL_MAX_DESCOMPRIMIDO = 20 * 1024 * 1024
#: Más filas descartadas que esto y el archivo no es una lista de correos: se corta antes
#: de armar una transacción con decenas de miles de filas de descarte.
MAX_DESCARTADOS = 20_000
ENCABEZADOS_EMAIL = frozenset({"email", "correo", "mail"})
LARGO_VALOR = 255

_validar_email = EmailValidator()


@dataclass
class LecturaExcel:
    """Lo que dejó el Excel: los correos a enviar y las filas que no."""

    #: ``[(fila, email)]`` en el orden del archivo, ya normalizados y sin repetir.
    validos: list = field(default_factory=list)
    #: ``[(fila, valor_leido, motivo)]`` con ``motivo`` en ``Descartado.Motivo``.
    descartados: list = field(default_factory=list)
    #: Filas con algún dato (sin contar el encabezado ni las filas vacías del todo).
    leidas: int = 0

    @property
    def invalidos(self):
        return sum(1 for _f, _v, m in self.descartados if m != Descartado.Motivo.DUPLICADO)

    @property
    def duplicados(self):
        return sum(1 for _f, _v, m in self.descartados if m == Descartado.Motivo.DUPLICADO)


def _sin_acentos(texto):
    normal = unicodedata.normalize("NFKD", texto)
    return "".join(c for c in normal if not unicodedata.combining(c))


def _texto(valor):
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip()


def _es_encabezado(valor):
    return _sin_acentos(_texto(valor)).lower() in ENCABEZADOS_EMAIL


def clave_de_duplicado(email):
    """La clave con la que se compara un correo contra los demás (RN-007-04).

    Sin mayúsculas (``casefold``) y sin acentos (NFKD): es lo que considera igual la
    collation ``*_ci`` de MySQL y MariaDB sobre la que vive el único (campaña, correo). Si
    acá fueran distintos y en la base iguales, el ``bulk_create`` reventaría con
    ``IntegrityError`` en vez de mandar el segundo a «Duplicado».
    """
    return _sin_acentos(email).casefold()


def normalizar_email(valor):
    """El correo como se compara y se guarda: recortado y en minúsculas (RN-007-03)."""
    return _texto(valor).lower()


def email_valido(email):
    if not email or len(email) > 254:
        return False
    try:
        _validar_email(email)
    except ValidationError:
        return False
    return True


def _verificar_descomprimido(archivo):
    """Rechaza el .xlsx cuyo contenido declarado no entra en el techo (como SEC-31).

    Lee solo el índice central del zip: no descomprime nada. Un archivo que no es un
    zip no se decide acá; lo informa ``load_workbook`` con el mensaje de siempre.
    """
    posicion = archivo.tell() if hasattr(archivo, "tell") else 0
    try:
        archivo.seek(0)
        try:
            with zipfile.ZipFile(archivo) as zf:
                total = sum(info.file_size for info in zf.infolist())
        except (zipfile.BadZipFile, OSError, ValueError):
            return
    finally:
        archivo.seek(posicion)
    if total > EXCEL_MAX_DESCOMPRIMIDO:
        raise ValidationError(
            "El Excel ocupa demasiado al descomprimirse. Dejá solo la hoja con los correos y volvé a subirlo."
        )


def _demasiados_descartados():
    raise ValidationError(
        f"El Excel tiene más de {MAX_DESCARTADOS:,} filas que no se pueden usar. ".replace(",", ".")
        + "Revisá que sea la lista de correos y que la columna sea la correcta."
    )


def parsear_destinatarios(archivo, *, tope=TOPE_DESTINATARIOS):
    """Lee el Excel y devuelve una :class:`LecturaExcel`.

    Levanta ``ValidationError`` si el archivo no es un ``.xlsx`` legible, pesa de más,
    no deja ningún correo válido o supera el tope de destinatarios. El archivo queda
    posicionado al principio para que el ``FileField`` lo pueda guardar después.
    """
    nombre = getattr(archivo, "name", "") or ""
    if not nombre.lower().endswith(".xlsx"):
        raise ValidationError("La lista tiene que ser un archivo Excel .xlsx (no se aceptan .xls ni .csv).")
    if (getattr(archivo, "size", 0) or 0) > EXCEL_MAX_BYTES:
        raise ValidationError("El Excel no puede superar los 2 MB.")
    _verificar_descomprimido(archivo)

    from openpyxl import load_workbook

    try:
        archivo.seek(0)
        libro = load_workbook(archivo, read_only=True, data_only=True)
    except Exception as exc:  # openpyxl levanta variantes según el archivo
        raise ValidationError("No se pudo leer el archivo: no es un Excel .xlsx válido.") from exc

    lectura = LecturaExcel()
    vistos = set()
    try:
        if not libro.worksheets:
            raise ValidationError("El Excel no tiene hojas.")
        hoja = libro.worksheets[0]
        columna = 0
        primera = True
        for numero, fila in enumerate(hoja.iter_rows(values_only=True), start=1):
            celdas = tuple(fila or ())
            if not any(_texto(c) for c in celdas):
                continue  # fila vacía del todo: no cuenta como leída
            if primera:
                # El encabezado se busca en la primera fila con datos, aunque arriba
                # haya filas en blanco.
                primera = False
                encabezado = next((i for i, c in enumerate(celdas) if _es_encabezado(c)), None)
                if encabezado is not None:
                    columna = encabezado
                    continue
            lectura.leidas += 1
            crudo = celdas[columna] if columna < len(celdas) else None
            valor = _texto(crudo)[:LARGO_VALOR]
            if len(lectura.descartados) > MAX_DESCARTADOS:
                _demasiados_descartados()
            if not valor:
                lectura.descartados.append((numero, "", Descartado.Motivo.VACIO))
                continue
            email = normalizar_email(crudo)
            if not email_valido(email):
                # RN-007-03 y caso límite: «a@x.com; b@y.com» en una celda no se separa.
                lectura.descartados.append((numero, valor, Descartado.Motivo.INVALIDO))
                continue
            clave = clave_de_duplicado(email)
            if clave in vistos:
                lectura.descartados.append((numero, valor, Descartado.Motivo.DUPLICADO))
                continue
            vistos.add(clave)
            lectura.validos.append((numero, email))
        if len(lectura.descartados) > MAX_DESCARTADOS:
            _demasiados_descartados()
    finally:
        libro.close()
        archivo.seek(0)

    if not lectura.validos:
        raise ValidationError(
            "El Excel no tiene ningún correo válido. Se espera una dirección por fila en la columna "
            "«email» (o «correo», «mail»), o en la columna A."
        )
    if len(lectura.validos) > tope:
        raise ValidationError(
            f"La lista tiene {len(lectura.validos):,} correos válidos y el tope por campaña es "
            f"{tope:,}. Partila en más de una campaña.".replace(",", ".")
        )
    return lectura
