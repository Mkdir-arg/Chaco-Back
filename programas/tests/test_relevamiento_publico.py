"""Tests del relevamiento de tipo Formulario público — Fase 1 (#290/#291/#292,
análisis #289): modelo, gateo RBAC del backoffice y ciclo de vida.
"""

from datetime import date, timedelta
from io import StringIO

from django.contrib.auth.models import Group, Permission, User
from django.contrib.contenttypes.models import ContentType
from django.core.cache import cache
from django.core.exceptions import ValidationError
from django.core.management import call_command
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from core import rbac
from legajos.models import Ciudadano
from programas.forms import RelevamientoForm
from programas.management.commands.seed_becas import ROL_ADMIN, ROL_TERRITORIAL
from programas.models import (
    AsignacionTerritorial,
    Convocatoria,
    Formulario,
    ListaEspera,
    ProgramaSiis,
    Relevamiento,
    Segmento,
)
from programas.services.reportes_becas import reporte_cupos
from users.models import Capacidad

CAP_PUBLICO = "becas.relevamiento.publico"


def _celdas_del_xlsx(contenido):
    """Todos los valores de la primera hoja, como texto: el export es un libro."""
    from io import BytesIO

    from openpyxl import load_workbook

    hoja = load_workbook(BytesIO(contenido), read_only=True).worksheets[0]
    return {str(celda.value) for fila in hoja.iter_rows() for celda in fila if celda.value is not None}


def _dar_capacidad_publico(grupo):
    """Suma la capacidad del formulario público a un rol existente (mismo
    mecanismo que ``seed_becas``: Permission sobre el ancla Capacidad)."""
    ct = ContentType.objects.get_for_model(Capacidad)
    codename = rbac.codename_de(CAP_PUBLICO)
    perm, _ = Permission.objects.get_or_create(
        content_type=ct, codename=codename, defaults={"name": "Formulario público"}
    )
    grupo.permissions.add(perm)


class _BasePublicoTest(TestCase):
    def setUp(self):
        call_command("seed_becas", stdout=StringIO())
        self.segmento = Segmento.objects.create(nombre="Segmento P", cupo_maximo=100)
        self.convocatoria = Convocatoria.objects.create(
            nombre="Conv P",
            segmento=self.segmento,
            fecha_inicio=date(2026, 1, 1),
            fecha_fin=date(2026, 12, 31),
        )
        self.territorial = User.objects.create_user("terri_p", password="x")
        self.territorial.groups.add(Group.objects.get(name=ROL_TERRITORIAL))
        AsignacionTerritorial.objects.create(segmento=self.segmento, territorial=self.territorial)

        # Admin del programa SIN la capacidad de público (estado post-deploy).
        self.admin = User.objects.create_user("admin_p", password="x")
        self.admin.groups.add(Group.objects.get(name=ROL_ADMIN))

        # Admin de otro grupo CON la capacidad (encendido vía Roles).
        self.admin_publico = User.objects.create_user("admin_pub", password="x")
        grupo_admin = Group.objects.get(name=ROL_ADMIN)
        self.admin_publico.groups.add(grupo_admin)
        _dar_capacidad_publico(grupo_admin)
        # Ojo: al compartir el grupo, self.admin también ganaría la capacidad.
        # Para el "sin capacidad" se usa un rol clonado sin ella.
        grupo_sin = Group.objects.create(name="Becas — Admin sin público")
        grupo_sin.permissions.set(grupo_admin.permissions.exclude(codename=rbac.codename_de(CAP_PUBLICO)))
        from users.models import RolMeta

        RolMeta.objects.create(grupo=grupo_sin, activo=True)
        self.admin.groups.set([grupo_sin])

    def _crear_publico(self, **extra):
        defaults = {
            "convocatoria": self.convocatoria,
            "tipo": Relevamiento.Tipo.PUBLICO,
            "fecha_asignada": date(2026, 6, 1),
            "fecha_hasta": date(2026, 6, 30),
        }
        defaults.update(extra)
        return Relevamiento.objects.create(**defaults)

    def _form_data(self, **extra):
        data = {
            "tipo": Relevamiento.Tipo.PUBLICO,
            "convocatoria": self.convocatoria.pk,
            "fecha_asignada": "2026-06-01T08:00",
            "fecha_hasta": "2026-06-30T18:00",
            "cupo_maximo": 50,
        }
        data.update(extra)
        return data


class ModeloPublicoTests(_BasePublicoTest):
    def test_alta_publico_token_en_curso_sin_territorial(self):
        rel = self._crear_publico()
        self.assertTrue(rel.es_publico)
        self.assertIsNotNone(rel.token_publico)
        self.assertEqual(rel.estado, Relevamiento.Estado.EN_CURSO)
        self.assertIsNone(rel.territorial)
        self.assertIn(str(rel.token_publico), rel.url_publica)
        self.assertNotIn(str(rel.pk), rel.url_publica.replace(str(rel.token_publico), ""))

    def test_tokens_unicos_en_altas_sucesivas(self):
        rel1 = self._crear_publico()
        rel2 = self._crear_publico()
        self.assertNotEqual(rel1.token_publico, rel2.token_publico)
        self.assertEqual(rel2.numero, rel1.numero + 1)

    def test_clean_rechaza_publico_con_territorial(self):
        rel = Relevamiento(
            convocatoria=self.convocatoria,
            tipo=Relevamiento.Tipo.PUBLICO,
            territorial=self.territorial,
            fecha_asignada=date(2026, 6, 1),
            fecha_hasta=date(2026, 6, 30),
        )
        with self.assertRaises(ValidationError) as ctx:
            rel.clean()
        self.assertIn("territorial", ctx.exception.message_dict)

    def test_clean_rechaza_territorial_sin_territorial_o_zona(self):
        rel = Relevamiento(
            convocatoria=self.convocatoria,
            fecha_asignada=date(2026, 6, 1),
            fecha_hasta=date(2026, 6, 30),
        )
        with self.assertRaises(ValidationError) as ctx:
            rel.clean()
        self.assertIn("territorial", ctx.exception.message_dict)
        rel.territorial = self.territorial
        with self.assertRaises(ValidationError) as ctx:
            rel.clean()
        self.assertIn("zona", ctx.exception.message_dict)

    def test_los_territoriales_existentes_no_cambian(self):
        rel = Relevamiento.objects.create(
            convocatoria=self.convocatoria,
            territorial=self.territorial,
            fecha_asignada=date(2026, 6, 1),
            zona="Zona X",
        )
        self.assertEqual(rel.tipo, Relevamiento.Tipo.TERRITORIAL)
        self.assertIsNone(rel.token_publico)
        self.assertEqual(rel.estado, Relevamiento.Estado.ASIGNADO)


class FormPublicoTests(_BasePublicoTest):
    def test_alta_publica_valida_sin_territorial_ni_zona(self):
        form = RelevamientoForm(data=self._form_data(), puede_publico=True)
        self.assertTrue(form.is_valid(), form.errors)
        rel = form.save()
        self.assertTrue(rel.es_publico)
        self.assertEqual(rel.estado, Relevamiento.Estado.EN_CURSO)
        self.assertEqual(rel.zona, "")
        self.assertEqual(rel.cupo_maximo, 50)

    def test_post_publico_con_territorial_se_rechaza(self):
        form = RelevamientoForm(data=self._form_data(territorial=self.territorial.pk), puede_publico=True)
        self.assertFalse(form.is_valid())
        self.assertIn("territorial", form.errors)

    def test_post_publico_sin_capacidad_se_rechaza(self):
        form = RelevamientoForm(data=self._form_data(), puede_publico=False)
        self.assertFalse(form.is_valid())
        self.assertIn("__all__", form.errors)

    def test_sin_capacidad_el_form_no_ofrece_tipo(self):
        form = RelevamientoForm(puede_publico=False)
        self.assertNotIn("tipo", form.fields)
        self.assertNotIn("padron", form.fields)

    def test_sin_capacidad_el_form_igual_ofrece_el_toggle_de_correo(self):
        """Cambio 44: los avisos por correo dejaron de ser del link público."""
        form = RelevamientoForm(puede_publico=False)
        self.assertIn("confirmar_por_email", form.fields)

    def test_flujo_territorial_sigue_exigiendo_territorial(self):
        data = self._form_data(tipo=Relevamiento.Tipo.TERRITORIAL)
        form = RelevamientoForm(data=data, puede_publico=True)
        self.assertFalse(form.is_valid())
        self.assertIn("territorial", form.errors)


class VistasPublicoTests(_BasePublicoTest):
    def test_crear_publico_redirige_al_detalle(self):
        self.client.force_login(self.admin_publico)
        resp = self.client.post(reverse("becas:relevamiento_crear"), self._form_data())
        rel = Relevamiento.objects.get(tipo=Relevamiento.Tipo.PUBLICO)
        self.assertRedirects(resp, reverse("becas:relevamiento_detalle", args=[rel.pk]))

    def test_detalle_muestra_link_copiable(self):
        rel = self._crear_publico()
        self.client.force_login(self.admin_publico)
        resp = self.client.get(reverse("becas:relevamiento_detalle", args=[rel.pk]))
        self.assertContains(resp, rel.url_publica)
        self.assertContains(resp, "data-copy-link")

    def test_sin_capacidad_no_ve_publicos_en_listado(self):
        rel = self._crear_publico()
        self.client.force_login(self.admin)
        resp = self.client.get(reverse("becas:relevamientos"))
        self.assertNotContains(resp, rel.nombre)
        self.client.force_login(self.admin_publico)
        resp = self.client.get(reverse("becas:relevamientos"))
        self.assertContains(resp, rel.nombre)
        self.assertContains(resp, "Formulario público")

    def test_sin_capacidad_no_ve_selector_de_tipo(self):
        self.client.force_login(self.admin)
        resp = self.client.get(reverse("becas:relevamientos"))
        self.assertNotContains(resp, "Tipo de relevamiento")

    def test_sin_capacidad_detalle_de_publico_403(self):
        rel = self._crear_publico()
        self.client.force_login(self.admin)
        resp = self.client.get(reverse("becas:relevamiento_detalle", args=[rel.pk]))
        self.assertEqual(resp.status_code, 403)

    def test_post_publico_sin_capacidad_no_crea(self):
        self.client.force_login(self.admin)
        self.client.post(reverse("becas:relevamiento_crear"), self._form_data())
        self.assertFalse(Relevamiento.objects.filter(tipo=Relevamiento.Tipo.PUBLICO).exists())

    def test_superusuario_puede_crear_publico(self):
        superuser = User.objects.create_superuser("root", "r@x.com", "x")
        self.client.force_login(superuser)
        resp = self.client.post(reverse("becas:relevamiento_crear"), self._form_data())
        self.assertEqual(resp.status_code, 302)
        self.assertTrue(Relevamiento.objects.filter(tipo=Relevamiento.Tipo.PUBLICO).exists())

    def test_reasignar_sobre_publico_no_aplica(self):
        rel = self._crear_publico()
        self.client.force_login(self.admin_publico)
        self.client.post(
            reverse("becas:relevamiento_reasignar", args=[rel.pk]),
            {"territorial": self.territorial.pk},
        )
        rel.refresh_from_db()
        self.assertIsNone(rel.territorial)

    def test_export_relevamientos_no_rompe_con_publico(self):
        self._crear_publico()
        self.client.force_login(self.admin_publico)
        resp = self.client.get(reverse("becas:convocatoria_export_relevamientos", args=[self.convocatoria.pk]))
        self.assertEqual(resp.status_code, 200)
        self.assertIn("Formulario público", resp.content.decode("utf-8-sig"))


class RnP13FueraDeLaPantallaTests(_BasePublicoTest):
    """SEC-22: RN-P13 valía en las pantallas y no en los datos que salen del sistema.

    Un rol con `becas.reportes.exportar` y sin `becas.relevamiento.publico` se bajaba el
    XLSX de respuestas —DNI, celular, email, GPS y todas las respuestas— de los casos
    del link público que en pantalla no podía ver, y los contaba en el cupo y en los
    reportes. El alcance del dashboard, de los reportes y del cupo ahora pasa por el
    mismo filtro que los listados.
    """

    def setUp(self):
        super().setUp()
        # El tablero se cachea por una clave de **proceso** que sobrevive al rollback
        # entre tests: sin esto, una entrada de otro test con el mismo pk contesta acá.
        cache.clear()
        self.programa_siis = ProgramaSiis.objects.create(nombre="Becas RN-P13", siis_programa_id=9501)
        self.segmento.programa = self.programa_siis
        self.segmento.save(update_fields=["programa", "modificado"])
        self.caso_publico = self._caso(self._crear_publico(), "40111000", "Publi")
        self.caso_territorial = self._caso(self._relevamiento_territorial(), "40111001", "Terri")

    def _relevamiento_territorial(self):
        return Relevamiento.objects.create(
            convocatoria=self.convocatoria,
            territorial=self.territorial,
            fecha_asignada=date(2026, 6, 1),
            fecha_hasta=date(2026, 6, 30),
            zona="Zona T",
        )

    def _caso(self, relevamiento, dni, nombre):
        ciudadano = Ciudadano.objects.create(dni=dni, nombre=nombre, apellido="RNP13")
        return Formulario.objects.create(
            relevamiento=relevamiento,
            ciudadano=ciudadano,
            estado=Formulario.Estado.APROBADO,
            celular="3624000000",
        )

    def test_el_cupo_no_muestra_los_casos_del_link_publico(self):
        self.client.force_login(self.admin)

        respuesta = self.client.get(reverse("becas:cupo_segmento", args=[self.segmento.pk]))

        self.assertEqual(respuesta.status_code, 200)
        self.assertContains(respuesta, self.caso_territorial.ciudadano.dni)
        self.assertNotContains(respuesta, self.caso_publico.ciudadano.dni)

    def test_con_la_capacidad_el_cupo_los_sigue_mostrando(self):
        self.client.force_login(self.admin_publico)

        respuesta = self.client.get(reverse("becas:cupo_segmento", args=[self.segmento.pk]))

        self.assertContains(respuesta, self.caso_publico.ciudadano.dni)

    def test_el_reporte_de_beneficiarios_no_los_exporta(self):
        self.client.force_login(self.admin)

        respuesta = self.client.get(reverse("becas:reporte_exportar", args=["beneficiarios", "csv"]))

        self.assertEqual(respuesta.status_code, 200)
        contenido = respuesta.content.decode("utf-8-sig")
        self.assertIn(self.caso_territorial.ciudadano.dni, contenido)
        self.assertNotIn(self.caso_publico.ciudadano.dni, contenido)

    def test_el_xlsx_de_respuestas_no_los_exporta(self):
        self.client.force_login(self.admin)

        respuesta = self.client.get(
            reverse(
                "becas:programa_dashboard_respuestas_xlsx",
                args=[self.programa_siis.pk, self.convocatoria.pk],
            )
        )

        self.assertEqual(respuesta.status_code, 200)
        valores = _celdas_del_xlsx(respuesta.content)
        self.assertIn(self.caso_territorial.ciudadano.dni, valores)
        self.assertNotIn(self.caso_publico.ciudadano.dni, valores)

    def test_con_la_capacidad_el_xlsx_los_sigue_trayendo(self):
        self.client.force_login(self.admin_publico)

        respuesta = self.client.get(
            reverse(
                "becas:programa_dashboard_respuestas_xlsx",
                args=[self.programa_siis.pk, self.convocatoria.pk],
            )
        )

        self.assertIn(self.caso_publico.ciudadano.dni, _celdas_del_xlsx(respuesta.content))

    def test_el_dashboard_no_los_cuenta(self):
        self.client.force_login(self.admin)

        respuesta = self.client.get(reverse("becas:programa_dashboard_datos", args=[self.programa_siis.pk]))

        self.assertEqual(respuesta.status_code, 200)
        indicadores = respuesta.json()["datos"]["indicadores"]
        self.assertEqual(indicadores["formularios_recibidos"], 1)
        self.assertEqual(indicadores["relevamientos_publicos"], 0)

    def test_el_dashboard_de_quien_puede_los_cuenta(self):
        """Y no comparte la entrada de caché con el de al lado (la huella lleva RN-P13)."""
        self.client.force_login(self.admin)
        self.client.get(reverse("becas:programa_dashboard_datos", args=[self.programa_siis.pk]))
        self.client.force_login(self.admin_publico)

        datos = self.client.get(reverse("becas:programa_dashboard_datos", args=[self.programa_siis.pk])).json()

        self.assertEqual(datos["datos"]["indicadores"]["formularios_recibidos"], 2)
        self.assertEqual(datos["datos"]["indicadores"]["relevamientos_publicos"], 1)


class CupoYReporteCuentanLaCapacidadTests(_BasePublicoTest):
    """D-22 oculta **personas**, no descuenta capacidad (revisión de la ronda 1 del #626).

    Con RN-P13 aplicado también a los agregados, `reporte_cupos` decía 3/97 sobre un
    segmento donde la stat card de la pantalla de Cupo y la aprobación
    (`services.cupo.get_cupo_stats`, que cuenta el segmento entero) decían 10/90: el
    operador sin la capacidad leía 97 lugares libres y la aprobación número 91 se le iba
    a lista de espera sin explicación. Los agregados se cuentan sin el filtro; lo que
    sigue filtrado es todo queryset que **liste filas** de personas.
    """

    APROBADOS_PUBLICOS = 7
    APROBADOS_TERRITORIALES = 3

    def setUp(self):
        super().setUp()
        cache.clear()
        self.publico = self._crear_publico()
        self.territorial_rel = Relevamiento.objects.create(
            convocatoria=self.convocatoria,
            territorial=self.territorial,
            fecha_asignada=date(2026, 6, 1),
            fecha_hasta=date(2026, 6, 30),
            zona="Zona T",
        )
        self.dnis_publicos = [
            self._caso(self.publico, f"4511100{i}", Formulario.Estado.APROBADO).ciudadano.dni
            for i in range(self.APROBADOS_PUBLICOS)
        ]
        self.dnis_territoriales = [
            self._caso(self.territorial_rel, f"4522200{i}", Formulario.Estado.APROBADO).ciudadano.dni
            for i in range(self.APROBADOS_TERRITORIALES)
        ]
        # Uno en espera de cada canal: el agregado de «Lista de espera» del reporte
        # tampoco es una lista de personas.
        for relevamiento, dni in ((self.publico, "45333001"), (self.territorial_rel, "45333002")):
            ListaEspera.objects.create(
                formulario=self._caso(relevamiento, dni, Formulario.Estado.ENVIADO),
                segmento=self.segmento,
                posicion=1 if relevamiento is self.publico else 2,
            )

    def _caso(self, relevamiento, dni, estado):
        ciudadano = Ciudadano.objects.create(dni=dni, nombre=f"Caso {dni}", apellido="Cupo")
        return Formulario.objects.create(relevamiento=relevamiento, ciudadano=ciudadano, estado=estado)

    def _fila_del_reporte(self, user):
        reporte = reporte_cupos(user)
        return next(fila for fila in reporte.filas if fila[0] == self.segmento.nombre)

    def test_el_reporte_cuenta_el_cupo_que_ya_consumio_el_link_publico(self):
        fila = self._fila_del_reporte(self.admin)

        # (Segmento, Cupo máximo, Distribuido, Ocupado, Disponible, Lista de espera, …)
        self.assertEqual(fila[1], 100)
        self.assertEqual(fila[3], self.APROBADOS_PUBLICOS + self.APROBADOS_TERRITORIALES)
        self.assertEqual(fila[4], 90)
        self.assertEqual(fila[5], 2)

    def test_el_reporte_dice_lo_mismo_con_la_capacidad_y_sin_ella(self):
        """El cupo es una propiedad del segmento, no de quién lo mira."""
        self.assertEqual(self._fila_del_reporte(self.admin), self._fila_del_reporte(self.admin_publico))

    def test_la_pantalla_de_cupo_dice_lo_mismo_que_el_reporte(self):
        self.client.force_login(self.admin)

        respuesta = self.client.get(reverse("becas:cupo_segmento", args=[self.segmento.pk]))

        self.assertEqual(respuesta.status_code, 200)
        stats = respuesta.context["stats"]
        fila = self._fila_del_reporte(self.admin)
        self.assertEqual((stats["cupo_ocupado"], stats["cupo_disponible"]), (10, 90))
        self.assertEqual((fila[3], fila[4]), (stats["cupo_ocupado"], stats["cupo_disponible"]))

    def test_las_personas_del_link_publico_siguen_fuera_de_las_listas(self):
        self.client.force_login(self.admin)

        respuesta = self.client.get(reverse("becas:cupo_segmento", args=[self.segmento.pk]))

        self.assertEqual(respuesta.context["n_beneficiarios"], self.APROBADOS_TERRITORIALES)
        self.assertEqual(respuesta.context["n_lista_espera"], 1)
        for dni in self.dnis_territoriales:
            self.assertContains(respuesta, dni)
        for dni in self.dnis_publicos:
            self.assertNotContains(respuesta, dni)

    def test_el_reporte_de_beneficiarios_sigue_sin_listar_los_publicos(self):
        """El agregado los cuenta; el CSV, que es una lista de personas, no los trae."""
        self.client.force_login(self.admin)

        contenido = self.client.get(reverse("becas:reporte_exportar", args=["beneficiarios", "csv"])).content.decode(
            "utf-8-sig"
        )

        for dni in self.dnis_territoriales:
            self.assertIn(dni, contenido)
        for dni in self.dnis_publicos:
            self.assertNotIn(dni, contenido)


class ApiCampoPublicoTests(_BasePublicoTest):
    def test_api_de_campo_no_expone_publicos(self):
        rel_publico = self._crear_publico(fecha_asignada=timezone.now(), fecha_hasta=timezone.now() + timedelta(days=5))
        self.client.force_login(self.territorial)
        listado = self.client.get(reverse("becas_api:relevamiento-list"))
        self.assertEqual(listado.status_code, 200)
        nombres = (
            [r["nombre"] for r in listado.json()["results"]]
            if "results" in listado.json()
            else [r["nombre"] for r in listado.json()]
        )
        self.assertNotIn(rel_publico.nombre, nombres)
        detalle = self.client.get(reverse("becas_api:relevamiento-detail", args=[rel_publico.pk]))
        self.assertEqual(detalle.status_code, 404)


class VencimientoPublicoTests(_BasePublicoTest):
    def test_publico_vencido_pasa_a_revision(self):
        ayer = timezone.now() - timedelta(days=1)
        rel = self._crear_publico(fecha_asignada=ayer - timedelta(days=5), fecha_hasta=ayer)
        call_command("procesar_vencimientos", stdout=StringIO())
        rel.refresh_from_db()
        self.assertEqual(rel.estado, Relevamiento.Estado.EN_REVISION)
        self.assertIsNotNone(rel.fecha_finalizado)

    def test_publico_vigente_no_se_toca(self):
        rel = self._crear_publico(
            fecha_asignada=timezone.now() - timedelta(days=1),
            fecha_hasta=timezone.now() + timedelta(days=5),
        )
        call_command("procesar_vencimientos", stdout=StringIO())
        rel.refresh_from_db()
        self.assertEqual(rel.estado, Relevamiento.Estado.EN_CURSO)
