"""Exportaciones del dashboard de Becas: formato, costo y techos (Ola 4 PR 5).

Tres fichas de la auditoría de octubre de 2026:

- **PERF-03** — armar el XLSX de «respuestas por persona» con 20.000 casos × 34
  columnas son ~7 s de CPU con el GIL tomado dentro del request (banco
  ``mariadb:10.11``); el **mismo** contenido en CSV son ~0,7 s. El botón nuevo no
  puede cambiar nada de lo que se descarga: acá se comparan los dos archivos
  generados, celda por celda.
- **G1b-11** — la exportación pedía una consulta ``JSON_EXTRACT`` + ``GROUP BY``
  **por pregunta** sobre todo el recorte, cada vez y aunque el bloque pedido no las
  usara.
- **G1b-12** — el período personalizado no tenía techo (``desde=2000-01-01`` y
  ``hasta=3999-12-31`` armaba ~104.000 semanas y las cacheaba, y la variación contra
  el período anterior se iba antes del año 1 con ``OverflowError``) y ``?recalcular=1``
  no tenía freno.
"""

import csv
from datetime import timedelta
from io import BytesIO, StringIO
from unittest import mock

from django.core.cache import cache
from django.db import connection
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from programas.forms_reportes import DashboardBecasFiltroForm
from programas.models import Formulario, PreguntaGlobal, TipoCampo
from programas.services import dashboard_becas as svc
from programas.tests.test_dashboard_becas import HOY, VENTANA, DashboardBecasBase


class ExportRespuestasFormatoTests(DashboardBecasBase):
    """PERF-03: la misma base cruda, en XLSX o en CSV."""

    def _url(self, convocatoria, formato):
        return reverse("becas:programa_dashboard_respuestas", args=[self.programa.pk, convocatoria.pk, formato])

    @staticmethod
    def _filas_del_xlsx(contenido):
        libro = load_workbook(BytesIO(contenido), read_only=True)
        return [
            ["" if celda is None else str(celda) for celda in fila]
            for fila in libro["Respuestas"].iter_rows(values_only=True)
        ]

    @staticmethod
    def _filas_del_csv(contenido):
        return list(csv.reader(StringIO(contenido.decode("utf-8-sig"))))

    def _dos_casos(self):
        ciudadana = self._ciudadano("30111222")
        self._formulario(
            self.rel_propio,
            Formulario.Estado.APROBADO,
            creado=HOY - timedelta(days=3),
            ciudadano=ciudadana,
            validado_renaper=True,
            data={
                "globales": {
                    str(self.q_laboral.pk): "Trabaja",
                    str(self.q_transporte.pk): ["Colectivo", "Moto"],
                },
                "requisitos": {str(self.r_cursa.pk): "Sí"},
            },
        )
        self._formulario(
            self.rel_publico,
            creado=HOY - timedelta(days=1),
            datos_identificacion={"dni": "40999888", "nombre": "Juan", "apellido": "Pérez"},
            data={"globales": {str(self.q_laboral.pk): "Estudia"}, "requisitos": {}},
        )

    def test_el_csv_trae_exactamente_lo_mismo_que_el_xlsx(self):
        self._dos_casos()
        self.client.force_login(self.admin)

        xlsx = self.client.get(self._url(self.conv_propia, "xlsx"))
        csv_ = self.client.get(self._url(self.conv_propia, "csv"))

        self.assertEqual((xlsx.status_code, csv_.status_code), (200, 200))
        self.assertIn('attachment; filename="becas_respuestas_', csv_["Content-Disposition"])
        self.assertTrue(csv_["Content-Disposition"].endswith('.csv"'))
        self.assertEqual(csv_["Content-Type"], "text/csv; charset=utf-8")

        filas_xlsx = self._filas_del_xlsx(xlsx.content)
        filas_csv = self._filas_del_csv(csv_.content)
        # El alcance encabeza los dos archivos (RN-16) y después va una fila en blanco.
        self.assertTrue(filas_xlsx[0][0].startswith("Alcance: Convocatoria "))
        self.assertEqual(len(filas_csv), len(filas_xlsx))
        for i, (del_xlsx, del_csv) in enumerate(zip(filas_xlsx, filas_csv)):
            with self.subTest(fila=i):
                self.assertEqual(del_csv, del_xlsx)
        # Y lo que se compara no es un par de archivos vacíos.
        self.assertEqual(len(filas_xlsx) - 3, 2)
        self.assertIn("Situación laboral", filas_xlsx[2])
        self.assertIn("Colectivo | Moto", [celda for fila in filas_csv for celda in fila])

    def test_los_dos_formatos_neutralizan_las_formulas(self):
        """SEC-20 no se negocia por formato: ``celda_segura`` corre en los dos caminos."""
        self._formulario(
            self.rel_propio,
            creado=HOY - timedelta(days=2),
            data={"globales": {str(self.q_laboral.pk): "=1+1"}, "requisitos": {}},
        )
        self.client.force_login(self.admin)

        xlsx = self._filas_del_xlsx(self.client.get(self._url(self.conv_propia, "xlsx")).content)
        csv_ = self._filas_del_csv(self.client.get(self._url(self.conv_propia, "csv")).content)

        columna = xlsx[2].index("Situación laboral")
        self.assertIn("'=1+1", [fila[columna] for fila in xlsx[3:]])
        self.assertIn("'=1+1", [fila[columna] for fila in csv_[3:]])

    def test_formato_invalido_permisos_y_alcance_valen_igual_para_el_csv(self):
        self._dos_casos()
        self.client.force_login(self.admin)

        self.assertEqual(self.client.get(self._url(self.conv_propia, "pdf")).status_code, 400)
        # Convocatoria de otro programa: 404 también en CSV, nunca datos ajenos.
        self.assertEqual(self.client.get(self._url(self.conv_otro, "csv")).status_code, 404)
        self.client.force_login(self._usuario_sin_reportes())
        self.assertEqual(self.client.get(self._url(self.conv_propia, "csv")).status_code, 403)
        self.client.logout()
        self.assertEqual(self.client.get(self._url(self.conv_propia, "csv")).status_code, 302)

    def test_la_pantalla_ofrece_los_dos_formatos(self):
        self.client.force_login(self.admin)

        pantalla = self.client.get(reverse("becas:programa_detalle", args=[self.programa.pk]))

        self.assertContains(pantalla, 'id="dash-respuestas-formato"')
        self.assertContains(
            pantalla, reverse("becas:programa_dashboard_respuestas", args=[self.programa.pk, 0, "FORMATO"])
        )


class DistribucionesDelExportTests(DashboardBecasBase):
    """G1b-11: una consulta agrupada por pregunta, sobre todo el recorte, sin caché."""

    def setUp(self):
        super().setUp()
        self._formulario(
            self.rel_propio,
            creado=HOY - timedelta(days=1),
            data={"globales": {str(self.q_laboral.pk): "Trabaja"}, "requisitos": {}},
        )
        self.client.force_login(self.admin)

    def _url(self, formato, **params):
        base = reverse("becas:programa_dashboard_exportar", args=[self.programa.pk, formato])
        return base + ("?" + "&".join(f"{k}={v}" for k, v in params.items()) if params else "")

    def _preguntas_extra(self, cuantas):
        for i in range(cuantas):
            PreguntaGlobal.objects.create(
                texto=f"Extra {i}", tipo=TipoCampo.SELECTOR, opciones=["A", "B"], orden=50 + i
            )

    @staticmethod
    def _por_pregunta(capturadas):
        return [q["sql"] for q in capturadas if "JSON_EXTRACT" in q["sql"]]

    def test_el_csv_de_otro_bloque_no_calcula_ninguna_distribucion(self):
        with CaptureQueriesContext(connection) as capturadas:
            respuesta = self.client.get(self._url("csv", periodo="90", bloque="convocatorias"))

        self.assertEqual(respuesta.status_code, 200)
        self.assertEqual(self._por_pregunta(capturadas), [])

    def test_el_csv_de_otro_bloque_no_crece_con_la_cantidad_de_preguntas(self):
        # Una pasada de descarte: la primera petición autenticada del test paga el
        # UPDATE del control de sesión única y desbalancearía la comparación.
        self.client.get(self._url("csv", periodo="90", bloque="convocatorias"))
        cache.clear()

        with CaptureQueriesContext(connection) as con_tres:
            self.client.get(self._url("csv", periodo="90", bloque="convocatorias"))
        self._preguntas_extra(6)
        cache.clear()

        with CaptureQueriesContext(connection) as con_nueve:
            self.client.get(self._url("csv", periodo="90", bloque="convocatorias"))

        self.assertEqual(len(svc.preguntas_graficables(self.admin, self.programa)), 9)
        self.assertEqual(len(con_nueve), len(con_tres))

    def test_el_csv_del_bloque_de_respuestas_las_sigue_calculando(self):
        respuesta = self.client.get(self._url("csv", periodo="90", bloque="respuestas"))

        self.assertEqual(respuesta.status_code, 200)
        filas = list(csv.reader(StringIO(respuesta.content.decode("utf-8-sig"))))
        self.assertEqual(filas[2][0], "Pregunta")
        self.assertIn("Situación laboral", [f[0] for f in filas[3:]])

    def test_el_xlsx_no_recalcula_lo_que_ya_esta_en_la_cache(self):
        self.client.get(self._url("xlsx", periodo="90"))

        with CaptureQueriesContext(connection) as segunda:
            respuesta = self.client.get(self._url("xlsx", periodo="90"))

        self.assertEqual(respuesta.status_code, 200)
        self.assertEqual(self._por_pregunta(segunda), [])

    def test_el_export_aprovecha_la_pregunta_que_calculo_la_pantalla(self):
        """La pantalla y el export comparten la entrada por pregunta: la que el tablero
        acaba de dibujar no se vuelve a agrupar."""
        self.client.get(reverse("becas:programa_dashboard_datos", args=[self.programa.pk]) + "?periodo=90")

        with CaptureQueriesContext(connection) as capturadas:
            self.client.get(self._url("xlsx", periodo="90"))

        catalogo = svc.preguntas_graficables(self.admin, self.programa)
        self.assertEqual(len(self._por_pregunta(capturadas)), len(catalogo) - 1)

    def test_las_distribuciones_cacheadas_dan_lo_mismo_que_calcularlas(self):
        sin_cache = svc.distribuciones_respuestas(self.admin, self.programa, VENTANA)

        primera = svc.distribuciones_cacheadas(self.admin, self.programa, VENTANA)
        segunda = svc.distribuciones_cacheadas(self.admin, self.programa, VENTANA)

        self.assertTrue(sin_cache)
        self.assertEqual([d.to_dict() for d in primera], [d.to_dict() for d in sin_cache])
        self.assertEqual([d.to_dict() for d in segunda], [d.to_dict() for d in sin_cache])


class PeriodoYRecalculoTests(DashboardBecasBase):
    """G1b-12: período personalizado sin techo y ``?recalcular=1`` sin freno."""

    def _form(self, **datos):
        return DashboardBecasFiltroForm({"periodo": "custom", **datos}, user=self.admin, programa=self.programa)

    def _url(self, **params):
        base = reverse("becas:programa_dashboard_datos", args=[self.programa.pk])
        return base + ("?" + "&".join(f"{k}={v}" for k, v in params.items()) if params else "")

    def _url_export(self, **params):
        base = reverse("becas:programa_dashboard_exportar", args=[self.programa.pk, "csv"])
        return base + ("?" + "&".join(f"{k}={v}" for k, v in params.items()) if params else "")

    def test_el_periodo_personalizado_no_puede_terminar_despues_de_hoy(self):
        hoy = timezone.localdate()

        form = self._form(desde=hoy.isoformat(), hasta=(hoy + timedelta(days=1)).isoformat())

        self.assertFalse(form.is_valid())
        self.assertIn("no puede ser posterior a hoy", " ".join(form.non_field_errors()))

    def test_el_periodo_personalizado_tiene_techo_de_cinco_anios(self):
        hoy = timezone.localdate()
        al_limite = self._form(desde=hoy.replace(year=hoy.year - 5).isoformat(), hasta=hoy.isoformat())
        self.assertTrue(al_limite.is_valid(), al_limite.errors)

        form = self._form(desde=hoy.replace(year=hoy.year - 5, day=1).isoformat(), hasta=hoy.isoformat())

        self.assertFalse(form.is_valid())
        self.assertIn("5 años", " ".join(form.non_field_errors()))

    def test_una_ventana_imposible_da_400_y_no_un_500(self):
        """``desde=2000-01-01&hasta=3999-12-31`` armaba ~104.000 semanas y las cacheaba,
        y la variación contra el período anterior se iba antes del año 1 con
        ``OverflowError``: la pantalla mostraba «no se pudieron calcular las métricas»."""
        self.client.force_login(self.admin)

        respuesta = self.client.get(self._url(periodo="custom", desde="2000-01-01", hasta="3999-12-31"))

        self.assertEqual(respuesta.status_code, 400)
        self.assertTrue(respuesta.json()["errores"])

    def test_una_ventana_en_los_primeros_anios_tambien_da_400(self):
        """El guard de G1b-12 tenía su propio 500: ``_hace_anios`` atrapaba el
        ``ValueError`` del 29 de febrero, pero ``replace`` levanta el mismo error
        cuando el año destino cae debajo de ``MINYEAR`` y el ``except`` repetía la
        operación. ``hasta`` en los años 1 a 5 reventaba en los dos endpoints."""
        self.client.force_login(self.admin)

        for anio in range(1, 6):
            parametros = {"periodo": "custom", "desde": "0001-01-01", "hasta": f"{anio:04d}-01-01"}
            with self.subTest(anio=anio):
                datos = self.client.get(self._url(**parametros))
                export = self.client.get(self._url_export(**parametros))

                self.assertEqual(datos.status_code, 400)
                self.assertTrue(datos.json()["errores"])
                self.assertEqual(export.status_code, 400)
                self.assertTrue(export.content.strip())

    def test_el_anio_uno_contra_el_9999_da_400_y_no_un_500(self):
        self.client.force_login(self.admin)
        parametros = {"periodo": "custom", "desde": "0001-01-01", "hasta": "9999-12-31"}

        datos = self.client.get(self._url(**parametros))
        export = self.client.get(self._url_export(**parametros))

        self.assertEqual(datos.status_code, 400)
        self.assertTrue(datos.json()["errores"])
        self.assertEqual(export.status_code, 400)
        self.assertTrue(export.content.strip())

    def test_el_periodo_personalizado_sigue_valiendo_arriba_del_piso(self):
        """El piso corta ventanas imposibles, no recortes reales: una ventana de un día
        apenas arriba de ``ANIO_MINIMO`` sigue siendo válida."""
        piso = DashboardBecasFiltroForm.ANIO_MINIMO

        self.assertFalse(self._form(desde=f"{piso - 1:04d}-01-01", hasta=f"{piso - 1:04d}-01-01").is_valid())
        self.assertTrue(self._form(desde=f"{piso:04d}-01-01", hasta=f"{piso:04d}-01-01").is_valid())

    def test_recalcular_se_atiende_una_vez_por_ventana(self):
        self.client.force_login(self.admin)
        self.client.get(self._url(periodo="90"))
        self._formulario(self.rel_propio, creado=HOY - timedelta(days=1))

        primero = self.client.get(self._url(periodo="90", recalcular="1"))
        self._formulario(self.rel_propio, creado=HOY - timedelta(days=1))
        segundo = self.client.get(self._url(periodo="90", recalcular="1"))

        self.assertFalse(primero.json()["desde_cache"])
        self.assertTrue(segundo.json()["desde_cache"], "el segundo «Actualizar» seguido recalculó el tablero entero")
        self.assertEqual(
            primero.json()["datos"]["indicadores"]["formularios_recibidos"],
            segundo.json()["datos"]["indicadores"]["formularios_recibidos"],
        )

    def test_vencido_el_freno_recalcular_vuelve_a_valer(self):
        self.client.force_login(self.admin)
        self.client.get(self._url(periodo="90"))

        with mock.patch.object(svc, "RECALCULO_MINIMO", 0):
            self.client.get(self._url(periodo="90", recalcular="1"))
            self._formulario(self.rel_propio, creado=HOY - timedelta(days=1))
            segundo = self.client.get(self._url(periodo="90", recalcular="1"))

        self.assertFalse(segundo.json()["desde_cache"])
        self.assertEqual(segundo.json()["datos"]["indicadores"]["formularios_recibidos"], 1)

    def test_el_freno_no_mezcla_dos_recortes(self):
        self.client.force_login(self.admin)

        primero = self.client.get(self._url(periodo="90", recalcular="1"))
        otro = self.client.get(self._url(periodo="30", recalcular="1"))

        self.assertFalse(primero.json()["desde_cache"])
        self.assertFalse(otro.json()["desde_cache"], "el freno de un recorte no puede frenar a otro")
